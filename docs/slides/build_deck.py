"""Build docs/slides/deck.pptx from docs/slides/content.yaml.

    python docs/slides/build_deck.py
    python docs/slides/build_deck.py --content other.yaml --out other.pptx

Each entry in content.yaml is one slide. Every slide has a `layout` and `notes` (speaker notes).
Content slides also take `title` and `section` (shown in the footer band). Layouts:

    cover          title (list of lines), subtitle, people (name, srn), course
    bullets        bullets, optional panel (heading, text) on the right
    numbers_top    numbers (value, label) in a row, bullets below
    numbers_side   numbers stacked on the left, optional heading and bullets on the right
    steps          steps (title, text), drawn as numbered rounded blocks
    compare        cards (heading, value, label, text) side by side, optional bullets below
    figure         figure (path from the repo root) with optional numbers and bullets, beside it
                   (text: side) or under it (text: below). Optional crops: a list of
                   [left, top, right, bottom] fractions; each crop is shown side by side, so a
                   wide multi-panel figure can show only the panels the slide talks about
    marker         a clearly marked placeholder for slides still to be written
    thankyou       title, text, optional side_text

Slides have at most four bullets. Colours come from polibias.config (the team palette), headings
use Georgia and body text Calibri. Paths in content.yaml are relative to the repository root.
"""

from __future__ import annotations

import argparse
import io
import sys
from pathlib import Path

import yaml
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))

from polibias import config  # noqa: E402

SLIDES_DIR = Path(__file__).resolve().parent
CONTENT_YAML = SLIDES_DIR / "content.yaml"
DECK_PPTX = SLIDES_DIR / "deck.pptx"

HEAD_FONT = "Georgia"
BODY_FONT = "Calibri"
# Big numbers use the body font: Georgia's old-style digits make "0" read like the letter "o"
NUMBER_FONT = BODY_FONT
MAX_BULLETS = 4

# Slide geometry in inches (16:9)
SLIDE_W = 13.333
SLIDE_H = 7.5
MARGIN = 0.6
TITLE_Y = 0.45
CONTENT_TOP = 1.45
FOOTER_H = 0.55
FOOTER_Y = SLIDE_H - FOOTER_H
CONTENT_BOTTOM = FOOTER_Y - 0.35
CONTENT_W = SLIDE_W - 2 * MARGIN
GAP = 0.3


def rgb(hex_colour: str) -> RGBColor:
    return RGBColor.from_string(hex_colour.lstrip("#").upper())


MIDNIGHT = rgb(config.MIDNIGHT)
DEEP_TEAL = rgb(config.DEEP_TEAL)
OCEAN_TEAL = rgb(config.OCEAN_TEAL)
TURQUOISE = rgb(config.TURQUOISE)
GOLD = rgb(config.GOLD)
GOLD_DARK = rgb(config.GOLD_DARK)
LIGHT = rgb(config.LIGHT)
MUTED = rgb(config.MUTED)
# Slightly darker than LIGHT, for cards that sit on a light slide
CARD_LIGHT = rgb("#E4DFD2")


# ---------- low-level helpers ----------


def set_background(slide, colour: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = colour


def add_box(slide, x, y, w, h, colour, rounded=False, radius=0.08, name=None):
    """A filled rectangle with no outline and no shadow."""
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        shape.adjustments[0] = radius
    shape.fill.solid()
    shape.fill.fore_color.rgb = colour
    shape.line.fill.background()
    shape.shadow.inherit = False
    if name:
        shape.name = name
    return shape


def add_text(slide, x, y, w, h, name=None, anchor=MSO_ANCHOR.TOP):
    """An empty text box with no inner padding; returns its text frame."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tf


def style_run(run, size, colour, font=BODY_FONT, bold=False, italic=False, spacing=None):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = colour
    if spacing is not None:
        run.font._rPr.set("spc", str(int(spacing * 100)))  # letter spacing in points


def write(tf, text, size, colour, font=BODY_FONT, bold=False, italic=False, spacing=None,
          align=PP_ALIGN.LEFT, first=True, space_after=0):
    """Write one paragraph (the first paragraph of the frame when first=True)."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    run = p.add_run()
    run.text = str(text)
    style_run(run, size, colour, font, bold, italic, spacing)
    return p


def set_bullet(paragraph, colour: RGBColor, indent_in=0.3) -> None:
    """Give a paragraph a real bullet (not a typed character) with a hanging indent."""
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(indent_in)))
    pPr.set("indent", str(-Inches(indent_in)))
    for tag in ("a:buClr", "a:buFont", "a:buChar", "a:buNone"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    bu_clr = pPr.makeelement(qn("a:buClr"), {})
    srgb = bu_clr.makeelement(qn("a:srgbClr"), {"val": str(colour)})
    bu_clr.append(srgb)
    pPr.append(bu_clr)
    pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))


def add_bullets(slide, bullets, x, y, w, h, size=18, colour=MIDNIGHT, bullet_colour=TURQUOISE,
                name="Bullets", anchor=MSO_ANCHOR.TOP, space_after=14):
    if not bullets:
        return
    if len(bullets) > MAX_BULLETS:
        raise ValueError(f"at most {MAX_BULLETS} bullets per slide, got {len(bullets)}")
    tf = add_text(slide, x, y, w, h, name=name, anchor=anchor)
    for i, item in enumerate(bullets):
        p = write(tf, item, size, colour, first=(i == 0), space_after=space_after)
        p.line_spacing = 1.05
        set_bullet(p, bullet_colour)


def add_number_block(slide, x, y, w, h, value, label, value_size=36, name="Number"):
    """A rounded Deep Teal block with a big gold number and an off-white label."""
    add_box(slide, x, y, w, h, DEEP_TEAL, rounded=True, radius=0.1, name=f"{name} block")
    pad = 0.25
    tf = add_text(slide, x + pad, y + pad * 0.8, w - 2 * pad, h - 1.6 * pad,
                  name=f"{name} text", anchor=MSO_ANCHOR.MIDDLE)
    write(tf, value, value_size, GOLD, font=NUMBER_FONT, bold=True, space_after=4)
    write(tf, label, 14, LIGHT, first=False)


def add_title(slide, title: str) -> None:
    """Short uppercase, letter-spaced title, top left. The first word is bold, the rest light."""
    words = title.upper().split(" ", 1)
    tf = add_text(slide, MARGIN, TITLE_Y, CONTENT_W, 0.6, name="Title", anchor=MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    first = p.add_run()
    first.text = words[0]
    style_run(first, 26, MIDNIGHT, font=HEAD_FONT, bold=True, spacing=4)
    if len(words) > 1:
        rest = p.add_run()
        rest.text = " " + words[1]
        style_run(rest, 26, MIDNIGHT, font=HEAD_FONT, spacing=4)


def add_footer(slide, section: str, number: int) -> None:
    """Deep Teal band carrying the section name (left) and the slide number (right)."""
    add_box(slide, 0, FOOTER_Y, SLIDE_W, FOOTER_H, DEEP_TEAL, name="Footer band")
    tf = add_text(slide, MARGIN, FOOTER_Y, 8, FOOTER_H, name="Footer section",
                  anchor=MSO_ANCHOR.MIDDLE)
    write(tf, section.upper(), 11, LIGHT, spacing=2)
    tf = add_text(slide, SLIDE_W - MARGIN - 2, FOOTER_Y, 2, FOOTER_H, name="Slide number",
                  anchor=MSO_ANCHOR.MIDDLE)
    write(tf, f"{number:02d}", 11, GOLD, bold=True, spacing=2, align=PP_ALIGN.RIGHT)


def crop_images(path: Path, crops) -> list[io.BytesIO]:
    """Cut [left, top, right, bottom] fractions out of an image, as in-memory PNGs."""
    parts = []
    with Image.open(path) as img:
        for left, top, right, bottom in crops:
            box = (round(left * img.width), round(top * img.height), round(right * img.width),
                   round(bottom * img.height))
            buf = io.BytesIO()
            img.crop(box).save(buf, format="PNG")
            buf.seek(0)
            parts.append(buf)
    return parts


def add_picture_fit(slide, path: Path, x, y, w, h, align="left", crops=None, gap=0.3):
    """Insert an image (or crops of it, side by side) as large as fits in the box, keeping the
    aspect ratio. Returns the bottom edge in inches."""
    images = crop_images(path, crops) if crops else [path]
    sizes = []
    for image in images:
        with Image.open(image) as img:
            sizes.append((img.width, img.height))
        if hasattr(image, "seek"):
            image.seek(0)
    # All parts share one height; their widths follow from their aspect ratios.
    ratios = [pw / ph for pw, ph in sizes]
    gaps = gap * (len(images) - 1)
    pic_h = min(h, (w - gaps) / sum(ratios))
    total_w = pic_h * sum(ratios) + gaps
    left = x if align == "left" else x + (w - total_w) / 2
    top = y + (h - pic_h) / 2
    for i, (image, ratio) in enumerate(zip(images, ratios)):
        source = str(image) if isinstance(image, Path) else image
        pic = slide.shapes.add_picture(source, Inches(left), Inches(top), Inches(pic_h * ratio),
                                       Inches(pic_h))
        pic.name = f"Figure {path.stem}" + (f" part {i + 1}" if crops else "")
        left += pic_h * ratio + gap
    return top + pic_h


# ---------- layouts ----------


def layout_cover(slide, s):
    set_background(slide, LIGHT)
    panel_y, panel_h, panel_w = 1.2, 3.9, 9.6
    add_box(slide, 0, panel_y, panel_w, panel_h, MIDNIGHT, name="Cover panel")
    tf = add_text(slide, 1.0, panel_y + 0.55, panel_w - 1.6, 2.0, name="Cover title",
                  anchor=MSO_ANCHOR.TOP)
    lines = s["title"] if isinstance(s["title"], list) else [s["title"]]
    for i, line in enumerate(lines):
        write(tf, line.upper(), 38, LIGHT, font=HEAD_FONT, bold=(i == 0), spacing=5,
              first=(i == 0), space_after=2)
    if s.get("subtitle"):
        tf = add_text(slide, 1.0, panel_y + 2.55, panel_w - 1.8, 1.0, name="Cover subtitle")
        write(tf, s["subtitle"], 18, GOLD, italic=True)
    tf = add_text(slide, 1.0, panel_y + panel_h + 0.45, 11.0, 1.2, name="Cover people")
    for i, person in enumerate(s.get("people", [])):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(4)
        name = p.add_run()
        name.text = person["name"]
        style_run(name, 18, MIDNIGHT, bold=True)
        srn = p.add_run()
        srn.text = "    " + person["srn"]
        style_run(srn, 16, MUTED, spacing=1)
    if s.get("course"):
        tf = add_text(slide, 1.0, SLIDE_H - 0.85, 11.0, 0.4, name="Cover course")
        write(tf, s["course"].upper(), 12, DEEP_TEAL, bold=True, spacing=3)


def layout_bullets(slide, s):
    panel = s.get("panel")
    text_w = 7.2 if panel else CONTENT_W
    add_bullets(slide, s.get("bullets"), MARGIN, CONTENT_TOP + 0.1, text_w,
                CONTENT_BOTTOM - CONTENT_TOP - 0.1, size=20, anchor=MSO_ANCHOR.MIDDLE,
                space_after=20)
    if panel:
        x = MARGIN + text_w + 0.6
        w = SLIDE_W - MARGIN - x
        y, h = CONTENT_TOP + 0.1, CONTENT_BOTTOM - CONTENT_TOP - 0.1
        add_box(slide, x, y, w, h, DEEP_TEAL, rounded=True, radius=0.06, name="Panel")
        tf = add_text(slide, x + 0.45, y + 0.45, w - 0.9, h - 0.9, name="Panel text",
                      anchor=MSO_ANCHOR.MIDDLE)
        write(tf, panel["heading"].upper(), 13, GOLD, bold=True, spacing=3, space_after=12)
        write(tf, panel["text"], 24, LIGHT, font=HEAD_FONT, first=False)


def layout_numbers_top(slide, s):
    numbers = s["numbers"]
    n = len(numbers)
    w = (CONTENT_W - GAP * (n - 1)) / n
    h = 1.75
    for i, num in enumerate(numbers):
        add_number_block(slide, MARGIN + i * (w + GAP), CONTENT_TOP, w, h, num["value"],
                         num["label"], name=f"Number {i + 1}")
    top = CONTENT_TOP + h + 0.45
    add_bullets(slide, s.get("bullets"), MARGIN, top, CONTENT_W, CONTENT_BOTTOM - top, size=18)


def layout_numbers_side(slide, s):
    numbers = s["numbers"]
    col_w = 4.0
    n = len(numbers)
    h = (CONTENT_BOTTOM - CONTENT_TOP - GAP * (n - 1)) / n
    h = min(h, 2.0)
    for i, num in enumerate(numbers):
        add_number_block(slide, MARGIN, CONTENT_TOP + i * (h + GAP), col_w, h, num["value"],
                         num["label"], value_size=40, name=f"Number {i + 1}")
    x = MARGIN + col_w + 0.7
    w = SLIDE_W - MARGIN - x
    top = CONTENT_TOP
    if s.get("heading"):
        tf = add_text(slide, x, top, w, 0.4, name="Heading")
        write(tf, s["heading"].upper(), 14, DEEP_TEAL, bold=True, spacing=3)
        top += 0.6
    add_bullets(slide, s.get("bullets"), x, top, w, CONTENT_BOTTOM - top, size=19)
    if s.get("caption"):
        tf = add_text(slide, x, CONTENT_BOTTOM - 0.35, w, 0.35, name="Caption",
                      anchor=MSO_ANCHOR.BOTTOM)
        write(tf, s["caption"], 12, MUTED, italic=True)


def layout_steps(slide, s):
    steps = s["steps"]
    n = len(steps)
    w = (CONTENT_W - GAP * (n - 1)) / n
    pill_h, text_h = 2.3, 1.2
    top = CONTENT_TOP + (CONTENT_BOTTOM - CONTENT_TOP - pill_h - 0.3 - text_h) / 2
    shades = [TURQUOISE, OCEAN_TEAL, DEEP_TEAL, MIDNIGHT]
    for i, step in enumerate(steps):
        x = MARGIN + i * (w + GAP)
        y = top
        add_box(slide, x, y, w, pill_h, shades[i % len(shades)], rounded=True, radius=0.18,
                name=f"Step {i + 1} block")
        tf = add_text(slide, x + 0.35, y + 0.3, w - 0.7, pill_h - 0.6, name=f"Step {i + 1} label",
                      anchor=MSO_ANCHOR.BOTTOM)
        write(tf, f"{i + 1:02d}", 44, GOLD, font=NUMBER_FONT, bold=True, space_after=2)
        write(tf, step["title"].upper(), 15, LIGHT, bold=True, spacing=2, first=False)
        tf = add_text(slide, x + 0.05, y + pill_h + 0.3, w - 0.1, text_h, name=f"Step {i + 1} text")
        write(tf, step["text"], 17, MIDNIGHT)
    add_bullets(slide, s.get("bullets"), MARGIN, CONTENT_BOTTOM - 0.5, CONTENT_W, 0.5, size=16)


def layout_compare(slide, s):
    cards = s["cards"]
    n = len(cards)
    w = (CONTENT_W - 0.5 * (n - 1)) / n
    has_bullets = bool(s.get("bullets"))
    h = 3.8
    for i, card in enumerate(cards):
        x = MARGIN + i * (w + 0.5)
        dark = card.get("dark", i == n - 1)
        bg, fg, accent, sub = ((DEEP_TEAL, LIGHT, GOLD, LIGHT) if dark
                               else (CARD_LIGHT, MIDNIGHT, GOLD_DARK, MUTED))
        add_box(slide, x, CONTENT_TOP, w, h, bg, rounded=True, radius=0.06, name=f"Card {i + 1}")
        tf = add_text(slide, x + 0.45, CONTENT_TOP + 0.4, w - 0.9, h - 0.8,
                      name=f"Card {i + 1} text")
        write(tf, card["heading"].upper(), 14, fg, bold=True, spacing=3, space_after=14)
        write(tf, card["value"], 54, accent, font=NUMBER_FONT, bold=True, first=False,
              space_after=0)
        write(tf, card["label"], 15, sub, first=False, space_after=16)
        write(tf, card["text"], 18, fg, first=False)
    if has_bullets:
        add_bullets(slide, s["bullets"], MARGIN, CONTENT_TOP + h + 0.5, CONTENT_W, 0.5, size=18)


def layout_figure(slide, s):
    """Figure with optional numbers and bullets. text: side (right column) or below (a strip
    under a full-width figure, for very wide figures; up to three short bullets)."""
    if s.get("text", "side") == "below":
        strip_h = 1.05
        fig_h = CONTENT_BOTTOM - CONTENT_TOP + 0.25 - strip_h - GAP
        bottom = add_picture_fit(slide, ROOT / s["figure"], MARGIN, CONTENT_TOP - 0.25, CONTENT_W,
                                 fig_h, align="center", crops=s.get("crops"))
        y = bottom + GAP
        x = MARGIN
        for i, num in enumerate(s.get("numbers", [])):
            add_number_block(slide, x, y, 4.3, strip_h, num["value"], num["label"],
                             value_size=24, name=f"Number {i + 1}")
            x += 4.3 + 0.5
        add_bullets(slide, s.get("bullets"), x, y, SLIDE_W - MARGIN - x, strip_h, size=15,
                    anchor=MSO_ANCHOR.MIDDLE, space_after=4)
        return
    side_w = 3.3
    fig_w = CONTENT_W - side_w - 0.4
    add_picture_fit(slide, ROOT / s["figure"], MARGIN, CONTENT_TOP - 0.25, fig_w,
                    CONTENT_BOTTOM - CONTENT_TOP + 0.3, crops=s.get("crops"))
    x = SLIDE_W - MARGIN - side_w
    top = CONTENT_TOP
    for i, num in enumerate(s.get("numbers", [])):
        add_number_block(slide, x, top, side_w, 1.55, num["value"], num["label"], value_size=32,
                         name=f"Number {i + 1}")
        top += 1.55 + 0.35
    add_bullets(slide, s.get("bullets"), x, top, side_w, CONTENT_BOTTOM - top, size=16,
                anchor=MSO_ANCHOR.MIDDLE)


def layout_marker(slide, s):
    x, y = MARGIN, CONTENT_TOP
    w, h = CONTENT_W, CONTENT_BOTTOM - CONTENT_TOP
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w),
                                 Inches(h))
    box.name = "Placeholder outline"
    box.adjustments[0] = 0.05
    box.fill.background()
    box.line.color.rgb = OCEAN_TEAL
    box.line.width = Pt(2)
    box.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    box.shadow.inherit = False
    tf = add_text(slide, x + 0.8, y + 0.6, w - 1.6, h - 1.2, name="Placeholder text",
                  anchor=MSO_ANCHOR.MIDDLE)
    write(tf, s.get("heading", "Placeholder").upper(), 16, OCEAN_TEAL, bold=True, spacing=3,
          align=PP_ALIGN.CENTER, space_after=12)
    write(tf, s.get("text", ""), 20, MIDNIGHT, align=PP_ALIGN.CENTER, first=False)


def layout_thankyou(slide, s):
    set_background(slide, LIGHT)
    panel_w = 9.8
    add_box(slide, 0, 0, panel_w, SLIDE_H, MIDNIGHT, name="Thank you panel")
    tf = add_text(slide, 1.0, 2.2, panel_w - 2.0, 3.0, name="Thank you text",
                  anchor=MSO_ANCHOR.MIDDLE)
    write(tf, s.get("title", "Thank you").upper(), 44, LIGHT, font=HEAD_FONT, bold=True,
          spacing=6, align=PP_ALIGN.CENTER, space_after=18)
    if s.get("text"):
        write(tf, s["text"], 18, GOLD, align=PP_ALIGN.CENTER, first=False)
    if s.get("side_text"):
        tf = add_text(slide, panel_w + 0.4, SLIDE_H - 1.3, SLIDE_W - panel_w - 0.8, 0.8,
                      name="Side text", anchor=MSO_ANCHOR.BOTTOM)
        write(tf, s["side_text"].upper(), 13, MIDNIGHT, bold=True, spacing=3,
              align=PP_ALIGN.RIGHT)


LAYOUTS = {
    "cover": layout_cover,
    "bullets": layout_bullets,
    "numbers_top": layout_numbers_top,
    "numbers_side": layout_numbers_side,
    "steps": layout_steps,
    "compare": layout_compare,
    "figure": layout_figure,
    "marker": layout_marker,
    "thankyou": layout_thankyou,
}
FULL_BLEED = {"cover", "thankyou"}  # dark panel slides: no title bar, no footer band


def build(content_path: Path = CONTENT_YAML, out_path: Path = DECK_PPTX) -> Path:
    slides = yaml.safe_load(content_path.read_text(encoding="utf-8"))["slides"]
    prs = Presentation()
    prs.slide_width = Emu(Inches(SLIDE_W))
    prs.slide_height = Emu(Inches(SLIDE_H))
    blank = prs.slide_layouts[6]
    for number, s in enumerate(slides, start=1):
        layout = s["layout"]
        if layout not in LAYOUTS:
            raise ValueError(f"slide {number}: unknown layout {layout!r}")
        if not s.get("notes"):
            raise ValueError(f"slide {number}: speaker notes are required")
        slide = prs.slides.add_slide(blank)
        if layout not in FULL_BLEED:
            set_background(slide, LIGHT)
            add_title(slide, s["title"])
        LAYOUTS[layout](slide, s)
        if layout not in FULL_BLEED:
            add_footer(slide, s.get("section", ""), number)
        slide.notes_slide.notes_text_frame.text = " ".join(s["notes"].split())
    prs.core_properties.title = "Categorizing News Articles on Political Bias"
    prs.core_properties.author = "Team 16"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    return out_path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--content", type=Path, default=CONTENT_YAML)
    parser.add_argument("--out", type=Path, default=DECK_PPTX)
    args = parser.parse_args()
    out = build(args.content, args.out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
